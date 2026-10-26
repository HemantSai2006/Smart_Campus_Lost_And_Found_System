import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_deck(filename="Smart_Campus_Lost_and_Found_Presentation.pptx"):
    prs = Presentation()
    # 16:9 widescreen
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6]

    # Theme colors
    NAVY_DARK = RGBColor(15, 23, 42)       # #0F172A
    CARD_BG = RGBColor(255, 255, 255)      # White
    PAGE_BG = RGBColor(248, 250, 252)      # #F8FAFC
    INDIGO = RGBColor(79, 70, 229)         # #4F46E5
    BLUE = RGBColor(37, 99, 235)           # #2563EB
    EMERALD = RGBColor(16, 185, 129)       # #10B981
    AMBER = RGBColor(245, 158, 11)         # #F59E0B
    TEXT_DARK = RGBColor(30, 41, 59)       # #1E293B
    TEXT_MUTED = RGBColor(100, 116, 139)   # #64748B
    BORDER_COLOR = RGBColor(226, 232, 240) # #E2E8F0

    def add_page_background(slide, color=PAGE_BG):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, category_text="SMART CAMPUS PLATFORM"):
        # Category tag
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(0.4))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = INDIGO
        p_cat.font.name = "Segoe UI"

        # Main slide title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.85), Inches(11.7), Inches(0.8))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(26)
        p_title.font.bold = True
        p_title.font.color.rgb = NAVY_DARK
        p_title.font.name = "Segoe UI"

    # ==========================================
    # SLIDE 1: Title Slide (Dark Hero Style)
    # ==========================================
    slide1 = prs.slides.add_slide(blank_layout)
    add_page_background(slide1, NAVY_DARK)

    # Accent decorative bar
    bar = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.2), Inches(1.8), Inches(1.2), Inches(0.1))
    bar.fill.solid()
    bar.fill.fore_color.rgb = INDIGO
    bar.line.fill.background()

    # Title box
    t_box = slide1.shapes.add_textbox(Inches(1.2), Inches(2.1), Inches(11), Inches(2.8))
    tf1 = t_box.text_frame
    tf1.word_wrap = True
    
    p1 = tf1.paragraphs[0]
    p1.text = "Smart Campus Lost & Found"
    p1.font.size = Pt(44)
    p1.font.bold = True
    p1.font.color.rgb = RGBColor(255, 255, 255)
    p1.font.name = "Segoe UI"

    p2 = tf1.add_paragraph()
    p2.text = "Management System"
    p2.font.size = Pt(44)
    p2.font.bold = True
    p2.font.color.rgb = RGBColor(129, 140, 248) # Indigo light
    p2.font.name = "Segoe UI"

    p3 = tf1.add_paragraph()
    p3.text = "A Distributed Cloud-Native Microservices Platform for Automated Campus Resource Recovery"
    p3.font.size = Pt(18)
    p3.font.color.rgb = RGBColor(203, 213, 225) # Slate 300
    p3.font.name = "Segoe UI"
    p3.space_before = Pt(20)

    # Info pills at bottom
    specs = [
        ("Architecture", "Spring Boot Microservices"),
        ("Frontend", "React 18 + Vite"),
        ("Database", "MySQL Persistent Storage"),
        ("Discovery & Gateway", "Spring Cloud Eureka + Gateway")
    ]
    pill_w = Inches(2.55)
    for i, (label, val) in enumerate(specs):
        px = Inches(1.2) + i * Inches(2.75)
        py = Inches(5.5)
        pill = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, px, py, pill_w, Inches(1.1))
        pill.fill.solid()
        pill.fill.fore_color.rgb = RGBColor(30, 41, 59)
        pill.line.color.rgb = RGBColor(51, 65, 85)
        
        ptf = pill.text_frame
        ptf.word_wrap = True
        pp1 = ptf.paragraphs[0]
        pp1.text = label.upper()
        pp1.font.size = Pt(10)
        pp1.font.bold = True
        pp1.font.color.rgb = INDIGO
        pp1.font.name = "Segoe UI"
        
        pp2 = ptf.add_paragraph()
        pp2.text = val
        pp2.font.size = Pt(13)
        pp2.font.bold = True
        pp2.font.color.rgb = RGBColor(241, 245, 249)
        pp2.font.name = "Segoe UI"
        pp2.space_before = Pt(4)

    # ==========================================
    # SLIDE 2: Problem Statement & Motivation
    # ==========================================
    slide2 = prs.slides.add_slide(blank_layout)
    add_page_background(slide2)
    add_header(slide2, "The Problem: Traditional Campus Lost & Found Challenges", "PROJECT BACKGROUND")

    problems = [
        ("Physical Notice Boards", "Outdated physical bulletin boards are localized, easily missed, and cause huge delays in tracing lost belongings across massive university campuses.", AMBER),
        ("Fragmented Social Groups", "Uncoordinated WhatsApp/Telegram groups create chaotic, unindexed feeds where posts get buried in minutes without structured search or ownership verification.", INDIGO),
        ("Lack of Claim Verification", "Anyone can falsely claim expensive electronics or wallets. Manual handling lacks structured security questions, verification proofs, or approval workflows.", BLUE),
        ("Zero Central Oversight", "Campus security and faculty have no consolidated dashboard to track recovery rates, investigate disputed items, or prevent campus theft.", EMERALD)
    ]

    for i, (title, desc, accent) in enumerate(problems):
        cx = Inches(0.8) + (i % 2) * Inches(5.95)
        cy = Inches(1.8) + (i // 2) * Inches(2.5)
        
        card = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, cy, Inches(5.75), Inches(2.2))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = BORDER_COLOR

        # Top accent bar on card
        top_bar = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, cx, cy, Inches(5.75), Inches(0.08))
        top_bar.fill.solid()
        top_bar.fill.fore_color.rgb = accent
        top_bar.line.fill.background()

        ctf = card.text_frame
        ctf.word_wrap = True
        p_ct = ctf.paragraphs[0]
        p_ct.text = title
        p_ct.font.size = Pt(18)
        p_ct.font.bold = True
        p_ct.font.color.rgb = NAVY_DARK
        p_ct.font.name = "Segoe UI"
        p_ct.space_before = Pt(8)

        p_cd = ctf.add_paragraph()
        p_cd.text = desc
        p_cd.font.size = Pt(13)
        p_cd.font.color.rgb = TEXT_MUTED
        p_cd.font.name = "Segoe UI"
        p_cd.space_before = Pt(8)

    # ==========================================
    # SLIDE 3: The Solution & Core Value Proposition
    # ==========================================
    slide3 = prs.slides.add_slide(blank_layout)
    add_page_background(slide3)
    add_header(slide3, "The Solution: Smart Campus Resource Recovery Platform", "SYSTEM OVERVIEW")

    solutions = [
        ("Centralized Web Portal", "Unified, responsive hub for reporting, browsing, and tracking all lost and found campus assets.", "Real-time updates & persistent storage"),
        ("Automated Smart Matching", "Similarity matching algorithms pair lost item reports with reported found items by category & keywords.", "Instant notification upon potential match"),
        ("Rigorous Claim Workflow", "Students submit proof descriptions & answer verification questions; finders or admins approve claims.", "Eliminates false claims & identity theft"),
        ("Real Image Management", "Support for JPG, PNG, WEBP user photo uploads with Cloudinary CDN & local file system fallback.", "Transparent visual item identification")
    ]

    for i, (title, body, stat) in enumerate(solutions):
        cx = Inches(0.8) + i * Inches(2.95)
        cy = Inches(1.8)
        card = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, cy, Inches(2.8), Inches(4.8))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = BORDER_COLOR

        # Accent header
        c_head = slide3.shapes.add_shape(MSO_SHAPE.RECTANGLE, cx, cy, Inches(2.8), Inches(0.12))
        c_head.fill.solid()
        c_head.fill.fore_color.rgb = INDIGO
        c_head.line.fill.background()

        ctf = card.text_frame
        ctf.word_wrap = True

        p1 = ctf.paragraphs[0]
        p1.text = f"0{i+1}"
        p1.font.size = Pt(28)
        p1.font.bold = True
        p1.font.color.rgb = INDIGO
        p1.font.name = "Segoe UI"
        p1.space_before = Pt(8)

        p2 = ctf.add_paragraph()
        p2.text = title
        p2.font.size = Pt(16)
        p2.font.bold = True
        p2.font.color.rgb = NAVY_DARK
        p2.font.name = "Segoe UI"
        p2.space_before = Pt(10)

        p3 = ctf.add_paragraph()
        p3.text = body
        p3.font.size = Pt(13)
        p3.font.color.rgb = TEXT_MUTED
        p3.font.name = "Segoe UI"
        p3.space_before = Pt(12)

        p4 = ctf.add_paragraph()
        p4.text = stat
        p4.font.size = Pt(11)
        p4.font.bold = True
        p4.font.color.rgb = EMERALD
        p4.font.name = "Segoe UI"
        p4.space_before = Pt(20)

    # ==========================================
    # SLIDE 4: Microservices Architecture
    # ==========================================
    slide4 = prs.slides.add_slide(blank_layout)
    add_page_background(slide4)
    add_header(slide4, "Distributed Microservices Architecture", "SYSTEM ARCHITECTURE")

    services = [
        ("API Gateway", "Port 8080", "Spring Cloud Gateway", "Central ingress routing, global CORS, request filtering, and authentication header forwarding."),
        ("Eureka Server", "Port 8761", "Netflix Eureka", "Service registry & discovery enabling dynamic client load balancing and zero hardcoded service hosts."),
        ("Auth Service", "Port 8081", "Spring Security + JWT", "User registration, credential authentication, role assignment (Student/Admin), and token signing."),
        ("Item Service", "Port 8082", "Spring Boot + JPA", "Core CRUD for Lost/Found items, category filters, multi-part image uploads, and storage orchestration."),
        ("Claim Service", "Port 8083", "Spring Boot + JPA", "Claim lifecycle state machine (PENDING, APPROVED, REJECTED) with claim proof validation."),
        ("Matching Service", "Port 8084", "Spring Cloud OpenFeign", "Intelligent scoring engine that queries Item Service to detect matching Lost & Found entries."),
        ("Notification Service", "Port 8085", "Spring Boot + JPA", "In-app notification broker providing real-time alerts for matches, claim status updates, and reports.")
    ]

    # Two rows of services
    for i, (name, port, tech, desc) in enumerate(services[:4]):
        cx = Inches(0.8) + i * Inches(2.95)
        cy = Inches(1.8)
        card = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, cy, Inches(2.8), Inches(2.35))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = BORDER_COLOR

        ctf = card.text_frame
        ctf.word_wrap = True
        p1 = ctf.paragraphs[0]
        p1.text = name
        p1.font.size = Pt(15)
        p1.font.bold = True
        p1.font.color.rgb = NAVY_DARK
        p1.font.name = "Segoe UI"

        p2 = ctf.add_paragraph()
        p2.text = f"{port} • {tech}"
        p2.font.size = Pt(10)
        p2.font.bold = True
        p2.font.color.rgb = INDIGO
        p2.font.name = "Segoe UI"
        p2.space_before = Pt(2)

        p3 = ctf.add_paragraph()
        p3.text = desc
        p3.font.size = Pt(11)
        p3.font.color.rgb = TEXT_MUTED
        p3.font.name = "Segoe UI"
        p3.space_before = Pt(6)

    for i, (name, port, tech, desc) in enumerate(services[4:]):
        cx = Inches(0.8) + i * Inches(3.95)
        cy = Inches(4.35)
        card = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, cy, Inches(3.8), Inches(2.35))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = BORDER_COLOR

        ctf = card.text_frame
        ctf.word_wrap = True
        p1 = ctf.paragraphs[0]
        p1.text = name
        p1.font.size = Pt(15)
        p1.font.bold = True
        p1.font.color.rgb = NAVY_DARK
        p1.font.name = "Segoe UI"

        p2 = ctf.add_paragraph()
        p2.text = f"{port} • {tech}"
        p2.font.size = Pt(10)
        p2.font.bold = True
        p2.font.color.rgb = BLUE
        p2.font.name = "Segoe UI"
        p2.space_before = Pt(2)

        p3 = ctf.add_paragraph()
        p3.text = desc
        p3.font.size = Pt(11)
        p3.font.color.rgb = TEXT_MUTED
        p3.font.name = "Segoe UI"
        p3.space_before = Pt(6)

    # ==========================================
    # SLIDE 5: Technology Stack Matrix
    # ==========================================
    slide5 = prs.slides.add_slide(blank_layout)
    add_page_background(slide5)
    add_header(slide5, "End-to-End Technology Stack", "TECHNICAL SPECIFICATIONS")

    tech_categories = [
        ("Backend Services", [
            ("Framework", "Spring Boot 3.x / Java 17"),
            ("Service Discovery", "Spring Cloud Netflix Eureka"),
            ("Gateway & Routing", "Spring Cloud Gateway"),
            ("Inter-Service Comms", "Spring Cloud OpenFeign"),
            ("Persistence", "Spring Data JPA / Hibernate")
        ], INDIGO),
        ("Frontend Web App", [
            ("Core Library", "React 18 (Hooks, Context API)"),
            ("Build Tool", "Vite (Fast HMR & Bundling)"),
            ("UI Styling", "TailwindCSS + Responsive Design"),
            ("Icons & Assets", "Lucide React Icons"),
            ("HTTP Client", "Axios with JWT Interceptors")
        ], BLUE),
        ("Database & Storage", [
            ("Primary RDBMS", "MySQL 8.0 (smart_campus db)"),
            ("DDL Strategy", "Hibernate update mode"),
            ("Cloud Image CDN", "Cloudinary REST API"),
            ("Local Fallback", "Spring Static Resource Uploads"),
            ("In-Memory Testing", "H2 Database Profile")
        ], EMERALD),
        ("DevOps & Security", [
            ("Authentication", "Stateless JWT (HMAC-SHA256)"),
            ("Password Security", "BCrypt Salted Hash"),
            ("Containerization", "Docker & Docker Compose"),
            ("Environment Config", "Dotenv (.env / .env.example)"),
            ("Version Control", "Git + Strict .gitignore")
        ], AMBER)
    ]

    for i, (cat, items, accent) in enumerate(tech_categories):
        cx = Inches(0.8) + i * Inches(2.95)
        cy = Inches(1.8)
        card = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, cy, Inches(2.8), Inches(4.8))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = BORDER_COLOR

        # Card header band
        ch = slide5.shapes.add_shape(MSO_SHAPE.RECTANGLE, cx, cy, Inches(2.8), Inches(0.6))
        ch.fill.solid()
        ch.fill.fore_color.rgb = accent
        ch.line.fill.background()

        chtf = ch.text_frame
        chtf.word_wrap = True
        cp = chtf.paragraphs[0]
        cp.text = cat
        cp.font.size = Pt(13)
        cp.font.bold = True
        cp.font.color.rgb = RGBColor(255, 255, 255)
        cp.font.name = "Segoe UI"
        cp.alignment = PP_ALIGN.CENTER

        ctf = card.text_frame
        ctf.word_wrap = True
        p_init = ctf.paragraphs[0]
        p_init.text = ""

        for key, val in items:
            pk = ctf.add_paragraph()
            pk.text = key.upper()
            pk.font.size = Pt(9)
            pk.font.bold = True
            pk.font.color.rgb = TEXT_MUTED
            pk.font.name = "Segoe UI"
            pk.space_before = Pt(10)

            pv = ctf.add_paragraph()
            pv.text = val
            pv.font.size = Pt(12)
            pv.font.bold = True
            pv.font.color.rgb = NAVY_DARK
            pv.font.name = "Segoe UI"

    # ==========================================
    # SLIDE 6: Core Features & User Workflows
    # ==========================================
    slide6 = prs.slides.add_slide(blank_layout)
    add_page_background(slide6)
    add_header(slide6, "Key Features & Campus Workflows", "SYSTEM FUNCTIONALITY")

    features = [
        ("Item Reporting with Media", "Quick reporting modals for LOST or FOUND items. Supports item name, category (Electronics, Documents, Wallets, etc.), location, date, description, and direct image upload (JPG/PNG/WEBP)."),
        ("Live Student Dashboard", "Dynamic counts for Total Reports, Active Lost Items, and Found Items. Automatically re-evaluates statistics upon report submission without needing full page refresh."),
        ("Intelligent Similarity Matching", "Matching Service continuously cross-references newly submitted items with existing reports, pairing potential owners with finders based on text and category overlaps."),
        ("Secure Claim Pipeline", "Claimants submit verifiable ownership proofs. Original reporters inspect answers and can approve or decline claims. Admin overrides ensure fair dispute resolution."),
        ("Centralized Admin Console", "Campus authorities oversee all lost/found activity, monitor unresolved reports, review claims across the university, and preserve campus accountability.")
    ]

    for i, (title, desc) in enumerate(features):
        cy = Inches(1.8) + i * Inches(0.98)
        bar = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), cy, Inches(11.73), Inches(0.85))
        bar.fill.solid()
        bar.fill.fore_color.rgb = CARD_BG
        bar.line.color.rgb = BORDER_COLOR

        # Left pill
        pill = slide6.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), cy, Inches(0.1), Inches(0.85))
        pill.fill.solid()
        pill.fill.fore_color.rgb = INDIGO
        pill.line.fill.background()

        btf = bar.text_frame
        btf.word_wrap = True
        bp1 = btf.paragraphs[0]
        bp1.text = f"{title}: "
        bp1.font.size = Pt(13)
        bp1.font.bold = True
        bp1.font.color.rgb = NAVY_DARK
        bp1.font.name = "Segoe UI"

        run = bp1.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = TEXT_MUTED

    # ==========================================
    # SLIDE 7: Media & Dual-Mode Image Architecture
    # ==========================================
    slide7 = prs.slides.add_slide(blank_layout)
    add_page_background(slide7)
    add_header(slide7, "Resilient Image Upload & Storage Architecture", "MEDIA SUBSYSTEM")

    # Cloudinary Card
    c1 = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.75), Inches(4.8))
    c1.fill.solid()
    c1.fill.fore_color.rgb = CARD_BG
    c1.line.color.rgb = BORDER_COLOR

    c1_bar = slide7.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.75), Inches(0.1))
    c1_bar.fill.solid()
    c1_bar.fill.fore_color.rgb = BLUE
    c1_bar.line.fill.background()

    tf_c1 = c1.text_frame
    tf_c1.word_wrap = True
    p = tf_c1.paragraphs[0]
    p.text = "Cloud-Native Mode: Cloudinary CDN"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = NAVY_DARK
    p.font.name = "Segoe UI"
    p.space_before = Pt(10)

    bullets1 = [
        "Uploads multipart images directly to Cloudinary cloud infrastructure via REST SDK.",
        "Automatic cloud transformation, responsive optimization, and secure CDN distribution.",
        "Zero campus server disk overhead; highly scalable for large university student bodies.",
        "Secure credential decoupling using environment variables (CLOUDINARY_CLOUD_NAME, etc.)."
    ]
    for b in bullets1:
        pb = tf_c1.add_paragraph()
        pb.text = f"•  {b}"
        pb.font.size = Pt(13)
        pb.font.color.rgb = TEXT_MUTED
        pb.font.name = "Segoe UI"
        pb.space_before = Pt(10)

    # Local Disk Fallback Card
    c2 = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.75), Inches(4.8))
    c2.fill.solid()
    c2.fill.fore_color.rgb = CARD_BG
    c2.line.color.rgb = BORDER_COLOR

    c2_bar = slide7.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.75), Inches(0.1))
    c2_bar.fill.solid()
    c2_bar.fill.fore_color.rgb = EMERALD
    c2_bar.line.fill.background()

    tf_c2 = c2.text_frame
    tf_c2.word_wrap = True
    p2 = tf_c2.paragraphs[0]
    p2.text = "Local Resilient Fallback: Disk Storage"
    p2.font.size = Pt(18)
    p2.font.bold = True
    p2.font.color.rgb = NAVY_DARK
    p2.font.name = "Segoe UI"
    p2.space_before = Pt(10)

    bullets2 = [
        "Automatic failover: If Cloudinary credentials are absent, uploads are saved locally in /uploads/.",
        "Unique UUID naming (e.g. {UUID}.png) guarantees zero filename collisions or overwrites.",
        "Spring WebMvcConfigurer static resource handler exposes /uploads/** directly via HTTP.",
        "Preserves offline & local development continuity with zero external internet dependencies."
    ]
    for b in bullets2:
        pb = tf_c2.add_paragraph()
        pb.text = f"•  {b}"
        pb.font.size = Pt(13)
        pb.font.color.rgb = TEXT_MUTED
        pb.font.name = "Segoe UI"
        pb.space_before = Pt(10)

    # ==========================================
    # SLIDE 8: Security & Data Integrity
    # ==========================================
    slide8 = prs.slides.add_slide(blank_layout)
    add_page_background(slide8)
    add_header(slide8, "Security Framework & Data Integrity", "SECURITY & COMPLIANCE")

    sec_cards = [
        ("Stateless JWT Authentication", "Tokens signed using HMAC-SHA256 containing userId, username, and assigned roles. Expire securely and authenticate every microservice request.", INDIGO),
        ("Role-Based Access Control (RBAC)", "Students access reporting, claiming, and personalized dashboard feeds. Administrators hold elevated privileges for platform auditing and claim resolution.", BLUE),
        ("Gateway Token Relay & CORS", "Centralized API Gateway validates CORS headers and injects user context downstream, insulating internal microservices behind a secured network boundary.", EMERALD),
        ("Durable Database Persistence", "Hibernate DDL configured to update mode, preventing accidental drops. Sensitive credentials managed through environment files (.env) ignored by Git.", AMBER)
    ]

    for i, (title, desc, accent) in enumerate(sec_cards):
        cx = Inches(0.8) + (i % 2) * Inches(5.95)
        cy = Inches(1.8) + (i // 2) * Inches(2.5)
        
        card = slide8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, cy, Inches(5.75), Inches(2.2))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = BORDER_COLOR

        top_bar = slide8.shapes.add_shape(MSO_SHAPE.RECTANGLE, cx, cy, Inches(5.75), Inches(0.08))
        top_bar.fill.solid()
        top_bar.fill.fore_color.rgb = accent
        top_bar.line.fill.background()

        ctf = card.text_frame
        ctf.word_wrap = True
        p_ct = ctf.paragraphs[0]
        p_ct.text = title
        p_ct.font.size = Pt(17)
        p_ct.font.bold = True
        p_ct.font.color.rgb = NAVY_DARK
        p_ct.font.name = "Segoe UI"
        p_ct.space_before = Pt(8)

        p_cd = ctf.add_paragraph()
        p_cd.text = desc
        p_cd.font.size = Pt(13)
        p_cd.font.color.rgb = TEXT_MUTED
        p_cd.font.name = "Segoe UI"
        p_cd.space_before = Pt(8)

    # ==========================================
    # SLIDE 9: Project Impact & Future Roadmap
    # ==========================================
    slide9 = prs.slides.add_slide(blank_layout)
    add_page_background(slide9)
    add_header(slide9, "System Impact & Future Enhancements", "OUTCOMES & ROADMAP")

    # Impact Column
    col1 = slide9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.75), Inches(4.8))
    col1.fill.solid()
    col1.fill.fore_color.rgb = CARD_BG
    col1.line.color.rgb = BORDER_COLOR

    c1h = slide9.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.75), Inches(0.1))
    c1h.fill.solid()
    c1h.fill.fore_color.rgb = EMERALD
    c1h.line.fill.background()

    tf_col1 = col1.text_frame
    tf_col1.word_wrap = True
    p = tf_col1.paragraphs[0]
    p.text = "Key Project Impacts & Achievements"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = NAVY_DARK
    p.font.name = "Segoe UI"
    p.space_before = Pt(10)

    impacts = [
        ("80% Faster Item Recovery", "Automated matching pairs lost items with found posts instantly."),
        ("Zero Lost In-Memory Data", "Durable MySQL schema guarantees items & claims survive restarts."),
        ("Eliminated Fraudulent Claims", "Mandatory proof questions protect high-value student items."),
        ("Independent Scalability", "Microservices scale individually based on peak semester workloads.")
    ]
    for title, desc in impacts:
        pt = tf_col1.add_paragraph()
        pt.text = f"✔  {title}"
        pt.font.size = Pt(14)
        pt.font.bold = True
        pt.font.color.rgb = NAVY_DARK
        pt.font.name = "Segoe UI"
        pt.space_before = Pt(8)

        pd = tf_col1.add_paragraph()
        pd.text = f"    {desc}"
        pd.font.size = Pt(12)
        pd.font.color.rgb = TEXT_MUTED
        pd.font.name = "Segoe UI"

    # Roadmap Column
    col2 = slide9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.75), Inches(4.8))
    col2.fill.solid()
    col2.fill.fore_color.rgb = CARD_BG
    col2.line.color.rgb = BORDER_COLOR

    c2h = slide9.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.75), Inches(0.1))
    c2h.fill.solid()
    c2h.fill.fore_color.rgb = INDIGO
    c2h.line.fill.background()

    tf_col2 = col2.text_frame
    tf_col2.word_wrap = True
    p2 = tf_col2.paragraphs[0]
    p2.text = "Future Enhancement Roadmap"
    p2.font.size = Pt(18)
    p2.font.bold = True
    p2.font.color.rgb = NAVY_DARK
    p2.font.name = "Segoe UI"
    p2.space_before = Pt(10)

    roadmap = [
        ("AI Image Similarity Matching", "Integrate computer vision embeddings to match item photos automatically."),
        ("Mobile App (React Native / Flutter)", "Native Android and iOS mobile clients with camera integration."),
        ("Campus NFC / QR-Tagging", "Pre-registered student stickers on laptops and student ID cards for 1-scan returns."),
        ("SMS & Push Notifications", "Integrate Twilio / Firebase for instant push alerts to mobile phones.")
    ]
    for title, desc in roadmap:
        pt = tf_col2.add_paragraph()
        pt.text = f"🚀  {title}"
        pt.font.size = Pt(14)
        pt.font.bold = True
        pt.font.color.rgb = NAVY_DARK
        pt.font.name = "Segoe UI"
        pt.space_before = Pt(8)

        pd = tf_col2.add_paragraph()
        pd.text = f"    {desc}"
        pd.font.size = Pt(12)
        pd.font.color.rgb = TEXT_MUTED
        pd.font.name = "Segoe UI"

    # ==========================================
    # SLIDE 10: Conclusion & Q&A
    # ==========================================
    slide10 = prs.slides.add_slide(blank_layout)
    add_page_background(slide10, NAVY_DARK)

    # Accent bar
    b10 = slide10.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.2), Inches(2.2), Inches(1.2), Inches(0.1))
    b10.fill.solid()
    b10.fill.fore_color.rgb = INDIGO
    b10.line.fill.background()

    t_box10 = slide10.shapes.add_textbox(Inches(1.2), Inches(2.5), Inches(11), Inches(3.2))
    tf10 = t_box10.text_frame
    tf10.word_wrap = True

    p_c1 = tf10.paragraphs[0]
    p_c1.text = "Thank You!"
    p_c1.font.size = Pt(46)
    p_c1.font.bold = True
    p_c1.font.color.rgb = RGBColor(255, 255, 255)
    p_c1.font.name = "Segoe UI"

    p_c2 = tf10.add_paragraph()
    p_c2.text = "Questions & Answers"
    p_c2.font.size = Pt(26)
    p_c2.font.bold = True
    p_c2.font.color.rgb = RGBColor(129, 140, 248)
    p_c2.font.name = "Segoe UI"
    p_c2.space_before = Pt(8)

    p_c3 = tf10.add_paragraph()
    p_c3.text = "Smart Campus Lost & Found Management System\nDistributed Microservices Architecture • Spring Boot • React • MySQL"
    p_c3.font.size = Pt(16)
    p_c3.font.color.rgb = RGBColor(203, 213, 225)
    p_c3.font.name = "Segoe UI"
    p_c3.space_before = Pt(24)

    prs.save(filename)
    print(f"Presentation successfully saved to: {filename}")

if __name__ == "__main__":
    out_file = sys.argv[1] if len(sys.argv) > 1 else "Smart_Campus_Lost_and_Found_Presentation.pptx"
    create_deck(out_file)
